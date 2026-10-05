"""Bulk issue import — migrates tickets from the legacy Cognito Forms
"RGMC IT Online Helpdesk" export (.xlsx / .csv) into the issues table.

The export schema is a flattened merge of two historical Cognito Forms
revisions (an older action/resolution form and a newer ticket form), so the
same logical field often appears twice under different headers (e.g. two
"Status" columns, two "Priority" columns, "Short Description" vs "Short
Description of Issue"). `_map_row` coalesces those into the single column
our schema actually uses.

Admin-only: this bulk-inserts directly into /issues, bypassing the normal
submission emails/bot notifications (nobody should get a "ticket submitted"
email for a ticket from months ago).
"""
import csv
import io
import re
from datetime import datetime, timezone

from flask import Blueprint, current_app, jsonify, request
from openpyxl import load_workbook

from services.guards import _require_admin
from services.supabase import supabase_req

issue_import_bp = Blueprint("issue_import", __name__)

BATCH_SIZE = 100
MAX_REPORTED_ROWS = 25

_STATUS_MAP = {
    "new": "open", "open": "open", "pending": "open", "unassigned": "open", "to do": "open",
    "assigned": "in_progress", "on-hold": "in_progress", "on hold": "in_progress", "onhold": "in_progress",
    "ongoing": "in_progress", "in progress": "in_progress", "reopened": "in_progress",
    "investigating": "in_progress", "in-progress": "in_progress",
    "resolved": "resolved", "completed": "resolved", "complete": "resolved", "fixed": "resolved",
    "closed": "closed", "done": "closed", "cancelled": "closed", "canceled": "closed",
    "rejected": "closed", "declined": "closed", "duplicate": "closed",
}

REQUIRED_LABELS = {
    "employee_name": "Name",
    "company_name":  "Company",
    "viber_number":  "Phone",
    "email":         "Email",
    "description":   "Description",
}

_DATE_FORMATS = (
    "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d",
    "%m/%d/%Y %H:%M:%S", "%m/%d/%Y %H:%M", "%m/%d/%Y %I:%M %p", "%m/%d/%Y",
)


def _norm_header(h) -> str:
    return re.sub(r"\s+", " ", str(h or "").replace("\xa0", " ")).strip().lower()


def _header_index_map(headers) -> dict:
    """normalized header text -> list of column indices, in file order
    (duplicate headers, e.g. the two "Status" columns, keep every index)."""
    out: dict[str, list[int]] = {}
    for i, h in enumerate(headers):
        key = _norm_header(h)
        if key:
            out.setdefault(key, []).append(i)
    return out


def _dedupe_header_labels(headers) -> list[str]:
    seen: dict[str, int] = {}
    out = []
    for h in headers:
        label = re.sub(r"\s+", " ", str(h or "").replace("\xa0", " ")).strip() or "(blank)"
        seen[label] = seen.get(label, 0) + 1
        out.append(label if seen[label] == 1 else f"{label} ({seen[label]})")
    return out


def _to_text(v):
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.isoformat()
    s = str(v).strip()
    return s or None


def _parse_date(v):
    if v is None:
        return None
    if isinstance(v, datetime):
        dt = v
    else:
        s = str(v).strip()
        if not s:
            return None
        dt = None
        for fmt in _DATE_FORMATS:
            try:
                dt = datetime.strptime(s, fmt)
                break
            except ValueError:
                continue
        if dt is None:
            try:
                dt = datetime.fromisoformat(s)
            except ValueError:
                return None
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=timezone.utc)
    return dt.isoformat()


def _map_priority(*raws):
    for raw in raws:
        if not raw:
            continue
        m = re.match(r"\s*p\s*([1-4])", str(raw), re.IGNORECASE)
        if m:
            return f"P{m.group(1)}"
    return None


def _map_status(*raws):
    """Returns (status, matched) — matched=False means every candidate
    value was blank/unrecognized and the 'open' default was used."""
    for raw in raws:
        key = _norm_header(raw)
        if key in _STATUS_MAP:
            return _STATUS_MAP[key], True
    return "open", False


def _read_rows(file_storage):
    filename = (file_storage.filename or "").lower()
    data = file_storage.read()

    if filename.endswith(".csv"):
        text = data.decode("utf-8-sig", errors="replace")
        rows = list(csv.reader(io.StringIO(text)))
        if not rows:
            return [], []
        return rows[0], rows[1:]

    if filename.endswith(".xlsx"):
        wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
        ws = wb.worksheets[0]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return [], []
        return list(rows[0]), [list(r) for r in rows[1:]]

    raise ValueError(
        "Unsupported file type — please export from Cognito Forms as .xlsx or .csv "
        "(legacy binary .xls files aren't supported; re-save as .xlsx first)."
    )


def _map_row(headers_idx: dict, dedup_labels: list[str], row: list, row_num: int):
    """Returns (issue_row | None, warnings, legacy_id, assignee_email).
    issue_row is None when a required field is missing — the row is skipped."""

    def col(name, which=0):
        idxs = headers_idx.get(name, [])
        if which >= len(idxs):
            return None
        idx = idxs[which]
        return row[idx] if idx < len(row) else None

    warnings = []

    legacy_id = _to_text(col("ticket number")) or _to_text(col("#"))
    name       = _to_text(col("name"))
    email      = _to_text(col("email"))
    phone      = _to_text(col("phone"))
    company    = _to_text(col("company"))
    department = _to_text(col("department")) or ""
    anydesk    = _to_text(col("anydesk id"))

    title = _to_text(col("short description of issue")) or _to_text(col("short description"))
    description = (
        _to_text(col("specific issues/concerns"))
        or _to_text(col("detailed description"))
        or title
    )

    request_category    = _to_text(col("request category"))
    request_subcategory = _to_text(col("request sub-category"))
    request_type_name   = _to_text(col("request type"))
    business_impact      = _to_text(col("business impact"))
    urgency               = _to_text(col("urgency"))
    ticket_type           = _to_text(col("ticket type"))
    software              = _to_text(col("software"))
    # "Software" is a vestige of the form's oldest revision and is almost
    # always a stale constant (e.g. every row says "Accounting") — prefer
    # the actual request category, which varies meaningfully per ticket.
    site_name = request_category or software or "General"

    priority = _map_priority(col("priority", 1), col("priority", 0))

    status, status_matched = _map_status(col("status", 1), col("status", 0))
    if not status_matched:
        raw_seen = _to_text(col("status", 1)) or _to_text(col("status", 0))
        warnings.append(f"Row {row_num}: status '{raw_seen or '(blank)'}' not recognized — defaulted to Open")

    created_at  = _parse_date(col("date submitted")) or _parse_date(col("date created"))
    resolved_at = _parse_date(col("date resolved")) or _parse_date(col("date completed"))

    assignee_name  = _to_text(col("assignee"))
    assignee_email = _to_text(col("assignee email"))

    resolution_notes = _to_text(col("action taken/ resolution"))
    remarks           = _to_text(col("remarks/comments"))
    if remarks:
        resolution_notes = f"{resolution_notes}\n\n{remarks}" if resolution_notes else remarks

    missing = [label for field, label in REQUIRED_LABELS.items()
               if not {"employee_name": name, "company_name": company, "viber_number": phone,
                        "email": email, "description": description}[field]]
    if missing:
        return None, [f"Row {row_num}: skipped — missing {', '.join(missing)}"], legacy_id, None

    raw_data = {label: _to_text(v) for label, v in zip(dedup_labels, row) if _to_text(v)}

    issue_row = {
        "site_name":             site_name,
        "employee_name":         name,
        "company_name":          company,
        "viber_number":          phone,
        "email":                 email,
        "department":            department,
        "description":           description,
        "title":                 title,
        "status":                status,
        "priority":              priority,
        "request_category":      request_category,
        "request_subcategory":   request_subcategory,
        "request_type_name":     request_type_name,
        "business_impact":       business_impact,
        "urgency":               urgency,
        "ticket_type":           ticket_type,
        "anydesk_id":            anydesk,
        "resolution_notes":      resolution_notes,
        "legacy_ticket_id":      legacy_id,
        "legacy_assignee_name":  assignee_name,
        "imported_from":         "cognito_forms",
        "imported_at":           datetime.now(timezone.utc).isoformat(),
        "legacy_raw_data":       raw_data,
    }
    if created_at:
        issue_row["created_at"] = created_at
    if resolved_at and status in ("resolved", "closed"):
        issue_row["resolved_at"] = resolved_at

    issue_row = {k: v for k, v in issue_row.items() if v not in (None, "")}
    return issue_row, warnings, legacy_id, assignee_email


@issue_import_bp.post("/api/admin/issues/import")
def import_issues():
    admin_username, err = _require_admin()
    if err:
        return jsonify(err[0]), err[1]

    f = request.files.get("file")
    if not f or not f.filename:
        return jsonify({"error": "No file provided"}), 400

    dry_run = request.form.get("dry_run", "true").strip().lower() != "false"

    try:
        headers, body_rows = _read_rows(f)
    except ValueError as exc:
        return jsonify({"error": str(exc)}), 400
    except Exception as exc:
        current_app.logger.error("issue import parse failed: %s", exc)
        return jsonify({"error": "Could not read this file. Make sure it's a valid .xlsx or .csv export."}), 400

    if not headers:
        return jsonify({"error": "The file appears to be empty."}), 400

    headers_idx  = _header_index_map(headers)
    dedup_labels = _dedupe_header_labels(headers)

    mapped = []       # (issue_row, legacy_id, assignee_email)
    skipped = []
    status_warnings = []
    total_rows = 0

    for i, row in enumerate(body_rows, start=2):
        if not any(c not in (None, "") for c in row):
            continue
        total_rows += 1
        issue_row, warnings, legacy_id, assignee_email = _map_row(headers_idx, dedup_labels, row, i)
        if issue_row is None:
            skipped.extend(warnings)
            continue
        status_warnings.extend(warnings)
        mapped.append((issue_row, legacy_id, assignee_email))

    # De-dupe against tickets already imported in a previous run
    legacy_ids = {lid for _, lid, _ in mapped if lid}
    existing_legacy_ids = set()
    if legacy_ids:
        try:
            esc = ",".join(f'"{lid}"' for lid in legacy_ids)
            existing = supabase_req("GET", "/issues", params={
                "legacy_ticket_id": f"in.({esc})",
                "select":           "legacy_ticket_id",
            })
            existing_legacy_ids = {r["legacy_ticket_id"] for r in (existing or []) if r.get("legacy_ticket_id")}
        except Exception as exc:
            current_app.logger.warning("issue import dedupe check failed: %s", exc)

    # Best-effort: match legacy "Assignee Email" to a real system username
    assignee_emails = {e.lower() for _, _, e in mapped if e}
    email_to_username = {}
    if assignee_emails:
        try:
            esc = ",".join(f'"{e}"' for e in assignee_emails)
            urows = supabase_req("GET", "/users", params={
                "email":  f"in.({esc})",
                "select": "username,email",
            })
            email_to_username = {u["email"].lower(): u["username"] for u in (urows or []) if u.get("email")}
        except Exception as exc:
            current_app.logger.warning("issue import assignee lookup failed: %s", exc)

    to_import = []
    duplicate_count = 0
    for issue_row, legacy_id, assignee_email in mapped:
        if legacy_id and legacy_id in existing_legacy_ids:
            duplicate_count += 1
            continue
        username = email_to_username.get((assignee_email or "").lower())
        if username:
            if issue_row["status"] in ("resolved", "closed"):
                issue_row["resolved_by"] = username
            issue_row["assigned_to"] = username
        to_import.append(issue_row)

    result = {
        "dry_run":          dry_run,
        "total_rows":       total_rows,
        "ready":            len(to_import),
        "duplicate":        duplicate_count,
        "skipped":          len(skipped),
        "skipped_reasons":  skipped[:MAX_REPORTED_ROWS],
        "status_warnings":  status_warnings[:MAX_REPORTED_ROWS],
        "status_warning_count": len(status_warnings),
    }

    if dry_run:
        result["preview"] = [{
            "legacy_ticket_id": r.get("legacy_ticket_id"),
            "title":            r.get("title") or (r.get("description") or "")[:80],
            "employee_name":    r.get("employee_name"),
            "company_name":     r.get("company_name"),
            "status":           r.get("status"),
            "priority":         r.get("priority"),
            "created_at":       r.get("created_at"),
        } for r in to_import[:10]]
        return jsonify(result)

    imported = 0
    failed_batches = 0
    for start in range(0, len(to_import), BATCH_SIZE):
        batch = to_import[start:start + BATCH_SIZE]
        try:
            rows = supabase_req("POST", "/issues", data=batch,
                                 extra_headers={"Prefer": "return=representation"})
            imported += len(rows or batch)
        except Exception as exc:
            failed_batches += 1
            current_app.logger.error("issue import batch failed (rows %s-%s): %s",
                                      start, start + len(batch), exc)

    result["imported"] = imported
    if failed_batches:
        result["error"] = (
            f"{failed_batches} batch(es) failed to save — {imported} of {len(to_import)} "
            "tickets were imported. Check the server logs and re-upload the same file; "
            "already-imported tickets will be skipped automatically."
        )
    current_app.logger.info("issue import by %s: %s/%s imported, %s duplicate, %s skipped",
                             admin_username, imported, len(to_import), duplicate_count, len(skipped))
    return jsonify(result)
